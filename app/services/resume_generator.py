"""Generate ATS-safe PDF (WeasyPrint) and DOCX from tailored content + master contact."""

from __future__ import annotations
from pathlib import Path
from datetime import datetime
from jinja2 import Environment, FileSystemLoader, select_autoescape
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.core.config import get_settings
from app.models.career import MasterCareer


def _templates_env():
    root = get_settings().project_root / "app" / "templates"
    return Environment(
        loader=FileSystemLoader(str(root)),
        autoescape=select_autoescape(["html", "xml"]),
    )


def render_html(contact: dict, tailored: dict) -> str:
    env = _templates_env()
    tpl = env.get_template("resume_ats.html")
    return tpl.render(contact=contact, r=tailored, generated_at=datetime.utcnow().isoformat())


def html_to_pdf(html: str, out_path: Path) -> Path:
    try:
        from weasyprint import HTML
        HTML(string=html).write_pdf(str(out_path))
    except Exception:
        # Fallback: write HTML so user can print-to-PDF
        out_path = out_path.with_suffix(".html")
        out_path.write_text(html, encoding="utf-8")
    return out_path


def generate_docx(contact: dict, tailored: dict, out_path: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    # Name
    p = doc.add_paragraph()
    run = p.add_run(contact.get("full_name", ""))
    run.bold = True
    run.font.size = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Contact line
    contact_bits = [
        contact.get("email"),
        contact.get("phone"),
        contact.get("location"),
        contact.get("linkedin"),
    ]
    contact_line = "  ·  ".join(b for b in contact_bits if b)
    p = doc.add_paragraph(contact_line)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(10)

    # Headline
    if tailored.get("headline"):
        p = doc.add_paragraph()
        run = p.add_run(tailored["headline"])
        run.bold = True
        run.font.size = Pt(12)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def section_heading(text: str):
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)

    # Summary
    if tailored.get("summary"):
        section_heading("Summary")
        doc.add_paragraph(tailored["summary"])

    # Skills
    skills = tailored.get("skills_order") or []
    if skills:
        section_heading("Skills")
        doc.add_paragraph(" · ".join(skills))

    # Experience
    exps = tailored.get("selected_experiences") or []
    if exps:
        section_heading("Experience")
        for exp in exps:
            p = doc.add_paragraph()
            run = p.add_run(f"{exp.get('title', '')}  |  {exp.get('company', '')}")
            run.bold = True
            if exp.get("dates"):
                p2 = doc.add_paragraph(exp["dates"])
                for r in p2.runs:
                    r.font.size = Pt(9)
                    r.italic = True
            for b in exp.get("bullets") or []:
                doc.add_paragraph(b, style="List Bullet")

    # Projects
    projs = tailored.get("selected_projects") or []
    if projs:
        section_heading("Projects")
        for proj in projs:
            p = doc.add_paragraph()
            run = p.add_run(proj.get("name", ""))
            run.bold = True
            if proj.get("description"):
                doc.add_paragraph(proj["description"])
            for o in proj.get("outcomes") or []:
                doc.add_paragraph(o, style="List Bullet")

    # Education
    edus = tailored.get("education") or []
    if edus:
        section_heading("Education")
        for edu in edus:
            line = f"{edu.get('degree', '')} — {edu.get('institution', '')}"
            if edu.get("end_year"):
                line += f" ({edu['end_year']})"
            doc.add_paragraph(line)

    # Certifications
    certs = tailored.get("certifications") or []
    if certs:
        section_heading("Certifications")
        for c in certs:
            doc.add_paragraph(f"{c.get('name', '')} — {c.get('issuer', '')}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def generate_documents(
    master: MasterCareer,
    tailored: dict,
    job_id: str,
    company: str,
) -> dict[str, str]:
    """Returns paths to generated files."""
    settings = get_settings()
    safe_company = "".join(c if c.isalnum() else "_" for c in company)[:40]
    stamp = datetime.utcnow().strftime("%Y%m%d_%H%M")
    base = settings.tailored_path / f"{job_id}_{safe_company}_{stamp}"
    base.parent.mkdir(parents=True, exist_ok=True)

    contact = master.contact.model_dump()
    html = render_html(contact, tailored)
    pdf_path = html_to_pdf(html, base.with_suffix(".pdf"))
    docx_path = generate_docx(contact, tailored, base.with_suffix(".docx"))

    return {
        "pdf": str(pdf_path.relative_to(settings.project_root)),
        "docx": str(docx_path.relative_to(settings.project_root)),
        "html": str(base.with_suffix(".html").relative_to(settings.project_root)) if pdf_path.suffix == ".html" else None,
    }
